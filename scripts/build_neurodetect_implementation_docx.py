from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = ROOT / "docs" / "neurodetect-implementation-report.md"
OUTPUT_DOCX = ROOT / "docs" / "NeuroDetect_Implementation_Report.docx"


def set_run_font(run, name: str, size: float, *, bold: bool = False, italic: bool = False, color: str = "000000") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        fmt = style.paragraph_format
        fmt.space_before = Pt(before)
        fmt.space_after = Pt(after)
        fmt.line_spacing = 1.10


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def configure_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run("NeuroDetect Implementation Report | ")
        set_run_font(run, "Calibri", 9, color="5D6C82")
        add_page_number(paragraph)


def clean_inline_markdown(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("`", "")
    text = text.replace("->", "->")
    return text


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(72)
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("NeuroDetect Implementation Report")
    set_run_font(run, "Calibri", 24, bold=True, color="15345D")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    run = p.add_run("Detailed project implementation record for the existing NeuroDetect research/demo prototype")
    set_run_font(run, "Calibri", 12, color="47698F")

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Document status", "Implemented and verified project summary"),
        ("Coverage", "Frontend, backend, database, Model 1, results integration, validation, and testing"),
        ("Project date", "August 23, 2026"),
        ("Important note", "NeuroDetect is a research/demo prototype and not a medical diagnostic system"),
    ]
    for row, (left, right) in zip(table.rows, rows):
        row.cells[0].text = left
        row.cells[1].text = right
    for row in table.rows:
        row.cells[0].width = Inches(1.9)
        row.cells[1].width = Inches(4.6)
        for idx, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    set_run_font(run, "Calibri", 10.5, bold=(idx == 0), color="142033")

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    note.paragraph_format.space_after = Pt(0)
    run = note.add_run(
        "This document summarizes the implementation work completed in the project repository, including the final "
        "assessment flow, mental health screening, cognitive tasks, Supabase persistence, Model 1 integration, the "
        "Results aggregation layer, and verification coverage."
    )
    set_run_font(run, "Calibri", 11, color="142033")

    doc.add_section(WD_SECTION.NEW_PAGE)


def add_code_block(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.right_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line.rstrip())
        set_run_font(run, "Consolas", 9.5, color="1F1F1F")


def add_paragraph_text(doc: Document, text: str, *, italic: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(clean_inline_markdown(text.strip()))
    set_run_font(run, "Calibri", 11, italic=italic, color="142033")


def build_docx() -> Path:
    source_text = SOURCE_MD.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_title_page(doc)

    lines = source_text.splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            code_lines: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(doc, code_lines)
            i += 1
            continue

        if stripped.startswith(">"):
            add_paragraph_text(doc, stripped.lstrip("> ").strip(), italic=True)
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            heading_text = clean_inline_markdown(stripped[level:].strip())
            if level == 1:
                p = doc.add_paragraph(style="Heading 1")
            elif level == 2:
                p = doc.add_paragraph(style="Heading 2")
            else:
                p = doc.add_paragraph(style="Heading 3")
            p.add_run(heading_text)
            i += 1
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(clean_inline_markdown(stripped[2:]))
            set_run_font(run, "Calibri", 11, color="142033")
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(clean_inline_markdown(re.sub(r"^\d+\.\s+", "", stripped)))
            set_run_font(run, "Calibri", 11, color="142033")
            i += 1
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            candidate = lines[i].strip()
            if not candidate:
                break
            if candidate.startswith(("```", "#", "> ", "- ")) or re.match(r"^\d+\.\s+", candidate):
                break
            paragraph_lines.append(candidate)
            i += 1
        add_paragraph_text(doc, " ".join(paragraph_lines))

    configure_footer(doc)
    doc.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


if __name__ == "__main__":
    path = build_docx()
    print(path)
